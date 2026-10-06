"""Generate editable widescreen PowerPoint and matching vector PDF sales decks."""
from pathlib import Path
from io import BytesIO
import json, math, hashlib, zipfile, datetime, re
from PIL import Image
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import ImageReader
from deck_content import DECKS, COMMON_SOURCES

ROOT=Path(__file__).resolve().parent
REPO=ROOT.parents[2]
W,H=960,540
DATE='05 OCTOBER 2026'
COMMIT='62fb640c'
FONTDIR=Path('C:/Windows/Fonts')
for name,file in [('Regular','segoeui.ttf'),('Bold','segoeuib.ttf'),('Light','segoeuil.ttf'),('Mono','consola.ttf')]:
    pdfmetrics.registerFont(TTFont(name,str(FONTDIR/file)))
PALETTES={
 'trade':dict(bg='#071E23',fg='#F3FAF7',muted='#B4C9C8',accent='#80EDBA',second='#FFC17A',panel='#113238',line='#2E5054',dark='#071E23'),
 'tech':dict(bg='#07111E',fg='#F5FAFF',muted='#AFC2D8',accent='#59D8FF',second='#FFA46B',panel='#102437',line='#284055',dark='#07111E'),
 'pulse':dict(bg='#F3F9FC',fg='#12364E',muted='#506C80',accent='#007EAF',second='#008579',panel='#E2F0F6',line='#BDD5E1',dark='#102E45'),
}

def rgb(h): return RGBColor.from_string(h.strip('#'))

class Deck:
    def __init__(self,data):
        self.d=data; self.p=PALETTES[data['theme']]; self.n=0; self.index=[]; self.boxes=[]; self.seller=[]
        self.stem=f"OperatorOS_{data['key']}_Sales_Deck"
        self.prs=Presentation(); self.prs.slide_width=Pt(W); self.prs.slide_height=Pt(H)
        self.prs.core_properties.title=f"OperatorOS | {data['key']} | Sales meeting"
        self.prs.core_properties.subject='A focused lane pitch with illustrative scenarios and a measurable pilot'
        self.prs.core_properties.author='Shotgun Ninjas Productions, LLC'
        self.prs.core_properties.keywords=f"OperatorOS, {data['key']}, sales, operations"
        self.c=canvas.Canvas(str(ROOT/(self.stem+'.pdf')),pagesize=(W,H),pageCompression=1)
        self.c.setTitle(self.prs.core_properties.title); self.c.setAuthor('Shotgun Ninjas Productions, LLC')
        self.c.setSubject(self.prs.core_properties.subject); self.c.setViewerPreference('DisplayDocTitle','true')
    def rect(self,x,y,w,h,fill=None,stroke=None,r=0):
        s=self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h))
        if r: s.adjustments[0]=min(.15,r/min(w,h))
        if fill: s.fill.solid(); s.fill.fore_color.rgb=rgb(fill)
        else: s.fill.background()
        if stroke: s.line.color.rgb=rgb(stroke); s.line.width=Pt(.8)
        else: s.line.fill.background()
        self.c.setFillColor(HexColor(fill or '#FFFFFF')); self.c.setStrokeColor(HexColor(stroke or '#FFFFFF')); self.c.setLineWidth(.8)
        if r: self.c.roundRect(x,H-y-h,w,h,r,stroke=bool(stroke),fill=bool(fill))
        else: self.c.rect(x,H-y-h,w,h,stroke=bool(stroke),fill=bool(fill))
        return s
    def line(self,x1,y1,x2,y2,color=None,width=1):
        color=color or self.p['line']
        s=self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(x1),Pt(y1),Pt(x2),Pt(y2))
        s.line.color.rgb=rgb(color); s.line.width=Pt(width)
        self.c.setStrokeColor(HexColor(color)); self.c.setLineWidth(width); self.c.line(x1,H-y1,x2,H-y2)
    def circle(self,x,y,diam,fill=None,stroke=None):
        s=self.slide.shapes.add_shape(MSO_SHAPE.OVAL,Pt(x),Pt(y),Pt(diam),Pt(diam))
        if fill: s.fill.solid(); s.fill.fore_color.rgb=rgb(fill)
        else: s.fill.background()
        if stroke:s.line.color.rgb=rgb(stroke);s.line.width=Pt(1)
        else:s.line.fill.background()
        self.c.setFillColor(HexColor(fill or '#FFFFFF'));self.c.setStrokeColor(HexColor(stroke or '#FFFFFF'));self.c.setLineWidth(1)
        self.c.circle(x+diam/2,H-y-diam/2,diam/2,fill=bool(fill),stroke=bool(stroke))
    def text(self,s,x,y,w,size=20,color=None,font='Regular',maxh=None,align='left',url=None):
        color=color or self.p['fg']; lines=[]
        for paragraph in s.split('\n'):
            acc=''
            for word in paragraph.split(' '):
                check=(acc+' '+word).strip()
                if acc and pdfmetrics.stringWidth(check,font,size)>w-2:
                    lines.append(acc); acc=word
                else: acc=check
            lines.append(acc)
        leading=size*1.19; h=leading*len(lines)+4
        if maxh and h>maxh:
            raise ValueError(f"Overflow {self.d['key']} {self.n} {s[:55]}: {h:.1f}>{maxh}")
        if y+h>H+1 or x+w>W+1: raise ValueError(f'Outside slide: {s}')
        box=self.slide.shapes.add_textbox(Pt(x),Pt(y),Pt(w+1),Pt(h))
        tf=box.text_frame; tf.clear();tf.word_wrap=False
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        tf.vertical_anchor=MSO_ANCHOR.TOP
        for i,ln in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=ln
            p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
            p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=Pt(leading)
            p.font.name='Consolas' if font=='Mono' else ('Segoe UI Light' if font=='Light' else 'Segoe UI')
            p.font.size=Pt(size);p.font.bold=font=='Bold';p.font.color.rgb=rgb(color)
            if url:
                for run in p.runs:run.hyperlink.address=url
            self.c.setFont(font,size);self.c.setFillColor(HexColor(color))
            xx=x if align=='left' else (x+w/2 if align=='center' else x+w)
            yy=H-y-pdfmetrics.getAscent(font,size)-i*leading
            if align=='left':self.c.drawString(xx,yy,ln)
            elif align=='center':self.c.drawCentredString(xx,yy,ln)
            else:self.c.drawRightString(xx,yy,ln)
        if url:self.c.linkURL(url,(x,H-y-h,x+w,H-y),relative=0,thickness=0)
        self.boxes.append(dict(slide=self.n,text=s,box=[x,y,w,h],size=size))
        return h
    def picture(self,path,x,y,w,h,cover=True):
        img=Image.open(path); iw,ih=img.size
        if cover:
            scale=max(w/iw,h/ih);sw,sh=iw*scale,ih*scale;cx=(sw-w)/2;cy=(sh-h)/2
            s=self.slide.shapes.add_picture(str(path),Pt(x),Pt(y),Pt(w),Pt(h))
            s.crop_left=s.crop_right=cx/sw;s.crop_top=s.crop_bottom=cy/sh
            self.c.saveState();p=self.c.beginPath();p.rect(x,H-y-h,w,h);self.c.clipPath(p,stroke=0,fill=0)
            self.c.drawImage(ImageReader(img),x-cx,H-y-h-cy,width=sw,height=sh,mask='auto');self.c.restoreState()
        else:
            scale=min(w/iw,h/ih);sw,sh=iw*scale,ih*scale
            self.slide.shapes.add_picture(str(path),Pt(x+(w-sw)/2),Pt(y+(h-sh)/2),Pt(sw),Pt(sh))
            self.c.drawImage(ImageReader(img),x+(w-sw)/2,H-y-(h+sh)/2,width=sw,height=sh,mask='auto')
    def brand(self,x=48,y=27,dark=False):
        self.picture(ROOT/'assets/operatoros-mark.png',x,y-4,30,30,False)
        self.text('OperatorOS',x+38,y,170,15,self.p['fg'] if not dark else '#F5FAFF','Bold')
    def new(self,title,kicker=None,notes='',source=None,cover=False):
        if self.n:self.c.showPage()
        self.n+=1;self.slide=self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.slide._element.cSld.set('name', title.replace('\n',' '))
        if self.n>=18:self.slide._element.set('show','0')
        self.rect(0,0,W,H,self.p['bg'])
        self.rect(0,0,7,H,self.p['accent'])
        self.c.bookmarkPage(f's{self.n}');self.c.addOutlineEntry(title.replace('\n',' '),f's{self.n}',0,False)
        if not cover:
            self.text(kicker or self.d['key'].upper(),48,31,725,10.5,self.p['accent'],'Bold')
            self.text(title,48,65,860,36,font='Bold',maxh=100)
            self.line(48,502,912,502)
            self.text('OPERATOROS  /  '+self.d['key'].upper(),48,515,440,9,self.p['muted'])
            self.text(f'{self.n:02d}  /  20',823,513,89,10,self.p['muted'],'Mono',align='right')
        src=source or COMMON_SOURCES
        notes=(notes+'\n\nEvidence basis: repository snapshot '+COMMIT+'; reviewed 2026-10-05. '+
            'Capabilities describe supported source workflows, not blanket acceptance of every target deployment or provider.\n'+
            '\n'.join(src)+'\nPublic lane: https://operatoros.net/for/'+self.d['lane']).strip()
        self.slide.notes_slide.notes_text_frame.text=notes
        self.index.append(dict(slide=self.n,title=title.replace('\n',' '),sources=src))
        self.seller.append(f'## {self.n:02d}. {title.replace(chr(10)," ")}\n\n{notes}\n')
    def label(self,s,x,y,w=400):self.text(s,x,y,w,10.5,self.p['accent'],'Bold')
    def band(self,s,y=444,h=42):
        self.rect(48,y,864,h,self.p['panel'],r=5);self.text(s,64,y+11,832,16,self.p['fg'],maxh=h-10)
    def save(self):
        self.c.showOutline();self.c.save();self.prs.save(ROOT/(self.stem+'.pptx'))
        (ROOT/(self.stem+'_Presenter_Notes.md')).write_text('# '+self.d['key']+' sales meeting guide\n\n'+ '\n'.join(self.seller),encoding='utf-8')
        (ROOT/(self.stem+'_Layout.json')).write_text(json.dumps(self.boxes,indent=2),encoding='utf-8')
        return dict(product=self.d['key'],slides=self.n,pdf=self.stem+'.pdf',pptx=self.stem+'.pptx',index=self.index)

def make(d):
    b=Deck(d); p=b.p
    # 01: Editorial hero with an unmodified local brand photo.
    b.new(d['title'],notes='OPEN: '+d['opener']+' Ask one discovery question before advancing: '+d['value_question']+' Position the product around an observable operating outcome. Cover imagery is conceptual brand art, not a customer or product screenshot.',cover=True)
    b.picture(ROOT/'assets'/d['image'],595,0,365,H)
    b.rect(577,0,18,H,p['accent']);b.brand()
    b.text(d['audience'],48,94,510,10.2,p['accent'],'Bold')
    b.text(d['key'],48,119,510,31,p['fg'],'Light')
    b.text(d['title'],48,178,520,44,p['fg'],'Bold',maxh=177)
    b.text(d['subtitle'],48,370,515,20,p['muted'],maxh=80)
    b.line(48,461,535,461,p['line']);b.text('A BETTER WAY TO OPERATE',48,478,350,10.5,p['accent'],'Bold')
    b.text('SHOTGUN NINJAS PRODUCTIONS  /  '+DATE,48,508,510,8.8,p['muted'])
    # 02: Pain and stakes.
    b.new(d['pain_title'],'THE FRICTION',notes='DISCOVER: '+d['value_question']+' Let the buyer identify the most expensive handoff. These are common scenario hypotheses, not measured customer results. Avoid assuming every pain applies.')
    for i,(head,body,end) in enumerate(d['pains']):
        x=48+i*296;b.text(f'0{i+1}',x,186,100,46,p['accent'],'Light')
        b.line(x,251,x+270,251);b.text(head,x,268,270,21,font='Bold',maxh=54)
        b.text(body,x,326,266,17.5,p['muted'],maxh=91)
        b.text(end,x,425,270,14.5,p['second'],'Bold',maxh=39)
    # 03: Value path.
    b.new(d['solution_title'],'THE CONNECTED WORKFLOW',notes=d['pain_close']+' Walk the buyer through a single record. These are workflow stages, not an assertion that the product performs every transition automatically.')
    for i,(head,body) in enumerate(d['workflow']):
        x=48+i*176;b.circle(x,205,33,p['accent']);b.text(str(i+1),x,211,33,15,p['bg'],'Bold',align='center')
        if i<4:b.line(x+40,222,x+166,222,p['line'],2)
        b.text(head,x,263,160,24,font='Bold');b.text(body,x,307,150,17,p['muted'])
    b.rect(48,382,864,84,p['panel'],r=5);b.label('THE RECORD YOU KEEP',65,396)
    b.text(d['solution_result'],65,420,830,18,maxh=45)
    # 04: OperatorOS explains the platform without suggesting all cores are included.
    b.new('Your core workflow.\nOne OperatorOS home base.','THE PLATFORM',notes='OperatorOS centrally manages sign-in, organizations, membership, roles, billing and application access. The current forward-sale package has one flagship per organization and one chosen eligible companion. Do not present all three flagship products as bundled into one subscription.')
    b.circle(76,209,214,None,p['line']);b.circle(96,229,174,p['panel'],p['accent'])
    b.picture(ROOT/'assets/operatoros-mark.png',155,253,58,58,False)
    b.text('OperatorOS',115,323,145,21,font='Bold',align='center')
    b.text('Your home base',105,355,163,13,p['muted'],align='center')
    for i,(a,c) in enumerate([('CORE WORKFLOW',d['key']),('CHOSEN COMPANION',d['companion']),('CENTRAL CONTROLS','People • roles • access • billing')]):
        y=193+i*84;b.line(291,316,391,y+34,p['line']);b.rect(391,y,521,69,p['panel'],r=5)
        b.label(a,411,y+11);b.text(c,411,y+30,476,22 if i<2 else 18,font='Bold')
    b.text('One flagship per organization. Expand with the eligible companions your team needs.',48,467,862,15,p['muted'])
    # 05: Six feature families.
    b.new(d['feature_title'],'WHAT THE TEAM GETS',notes='Use this as a capability map, not a rapid feature recital. Pick the two groups connected to the buyer’s pain, then continue to the workday view. Full feature checklist is in the appendix.')
    for i,(head,body) in enumerate(d['features']):
        x=48+(i%3)*296;y=191+(i//3)*146
        b.rect(x,y,272,129,p['panel'],r=6);b.text(f'0{i+1}',x+16,y+14,42,11,p['accent'],'Mono')
        b.text(head,x+16,y+34,240,19,font='Bold',maxh=49)
        b.text(body,x+16,y+66,238,15.5,p['muted'],maxh=61)
    # 06: Source-based differentiating workday brief.
    b.new(d['edge_title'],d['edge'],notes='This workday brief is grounded in persisted facts and links to existing records. '+d['edge_note']+' The drawing explains priority categories; it is not a screenshot or a claim about a real account.')
    b.text(d['edge_intro'],48,208,290,24,p['muted'],maxh=164)
    for i,(a,c) in enumerate(d['edge_rows']):
        y=195+i*60;b.rect(383,y,529,54,p['panel'],r=4);b.text(f'0{i+1}',397,y+16,35,14,p['accent'],'Mono')
        b.text(a,439,y+7,452,17,font='Bold');b.text(c,439,y+32,452,13.5,p['muted'])
    b.text(d['edge_result'],48,425,292,17,p['accent'],'Bold',maxh=48)
    b.text('WORKFLOW DIAGRAM / ILLUSTRATIVE',383,454,528,9.5,p['muted'])
    # 07-09: Three thoughtful example cases, each with tangible output and metric.
    for j,case in enumerate(d['cases']):
        b.new(case['title'],f"ILLUSTRATIVE SCENARIO {j+1:02d} / {case['tag']}",notes=case['talk']+' This is an invented, representative scenario, not a customer case study or a measured outcome. Trigger: '+case['trigger']+' Outcome to evaluate: '+case['measure'])
        b.rect(48,187,270,282,p['panel'],r=6);b.label('THE MOMENT',66,205,230)
        b.text(case['trigger'],66,234,234,22,maxh=197)
        for i,(a,c) in enumerate(case['steps']):
            y=187+i*67;b.circle(351,y+3,24,p['accent']);b.text(str(i+1),351,y+7,24,12,p['bg'],'Bold',align='center')
            if i<2:b.line(363,y+29,363,y+65,p['line'],1.5)
            b.text(a,391,y,510,18,font='Bold');b.text(c,391,y+25,510,15.7,p['muted'],maxh=42)
        b.label('WHAT THE TEAM CAN SHOW',350,396,540);b.text(case['leave'],350,415,560,18,maxh=48)
        b.text('MEASURE  /  '+case['measure'],48,477,864,10.5,p['muted'])
    # 10: Demo story, four actions and an outcome.
    b.new(d['demo_title'],'THE LIVE DEMO STORY',notes='Use a dedicated, approved demo workspace and synthetic records. Prepare the complete flow before the meeting. Do not attempt live payments, real messages, patient data or endpoint changes during the demo. If the tenant is not accepted, present this as a workflow discussion rather than imply live operation.')
    for i,(n,a,c) in enumerate(d['demo_steps']):
        x=48+i*222;b.rect(x,195,198,238,p['panel'],r=5)
        b.text(n,x+16,211,165,43,p['accent'],'Light');b.text(a,x+16,274,166,21,font='Bold',maxh=54)
        b.text(c,x+16,342,165,16,p['muted'],maxh=83)
    b.text('SUCCESS CHECK',48,457,135,10.5,p['accent'],'Bold');b.text(d['demo_proof'],200,451,712,17,maxh=45)
    # 11: Relevant ecosystem expansion, not an overwhelming catalog.
    b.new('Extend the workflow\nwhere it matters.','USEFUL COMPANIONS',notes=d['companion_note']+' Each companion still requires appropriate organization access. A useful combination does not imply an automatic cross-module connector. Core feature claims and companion readiness remain separate.')
    for i,(name,a,c,choice) in enumerate(d['companions']):
        x=48+i*444;b.rect(x,190,420,237,p['panel'],r=6);b.rect(x,190,420,4,p['accent'] if i==0 else p['second'])
        b.text(name,x+22,211,376,27,font='Bold');b.text(a,x+22,258,376,21,p['accent'],'Bold')
        b.text(c,x+22,302,372,18,p['muted'],maxh=85);b.text(choice,x+22,399,376,12.5,p['fg'],'Bold')
    b.text(d['companion_note'],48,446,864,12,p['muted'],maxh=47)
    # 12: Clear integration fit.
    b.new('Fit the workflow\nto your real environment.','WORK WITH THE TOOLS YOU RELY ON',notes='Clarify what the buyer wants to keep. Identify the required connection and its owner; verify before committing to an integration scope. Current source capabilities and public lane pages are the claim boundary.')
    for i,(tag,a,c) in enumerate(d['fit']):
        x=48+i*296;b.label(tag.upper(),x,199,270);b.line(x,229,x+272,229,p['accent'],2)
        b.text(a,x,255,270,26,font='Bold',maxh=70);b.text(c,x,350,270,20,p['muted'],maxh=120)
    # 13: Accountability or information control.
    b.new(d['trust_title'],'CONFIDENCE IN THE RECORD',notes=d['trust_footer']+' Separate product controls from team process. A recorded note is not proof of a real external action unless the appropriate supporting evidence is retained.')
    for i,(a,c) in enumerate(d['trust']):
        x=48+(i%2)*444;y=191+(i//2)*114;b.line(x,y,x+418,y,p['line'])
        b.text(a,x,y+14,415,22,font='Bold');b.text(c,x,y+50,404,17,p['muted'],maxh=51)
    b.text(d['trust_footer'],48,454,864,12.5,p['muted'],maxh=39)
    # 14: Buyer-defined value, no imaginary ROI.
    b.new('Measure the change\nin your own operation.','A BUSINESS CASE YOU CAN TEST',notes='Do not invent an ROI percentage. Define a baseline and an equivalent pilot sample, then review these measures. If monetizing time, use the buyer’s agreed hourly cost and distinguish released capacity from actual cash savings. Correlation during a pilot does not establish causation.')
    for i,(a,c) in enumerate(d['metrics']):
        x=48+(i%2)*444;y=190+(i//2)*101;b.rect(x,y,420,86,p['panel'],r=5)
        b.text(a,x+18,y+15,384,20,font='Bold');b.text(c,x+18,y+49,384,14.7,p['muted'],maxh=37)
    b.text(d['value_question'],48,409,864,25,p['accent'],'Bold',maxh=66)
    b.text('Agree the baseline, sample and success threshold together. No outcome guarantee is implied.',48,479,864,10.5,p['muted'])
    # 15: Offer structure, transparent live pricing limitation in notes.
    b.new('Start focused.\nGive the team room to work.','THE OPERATOROS APPLICATION STACK',notes='Verified source packaging: one flagship per organization, five seats and one eligible organization-wide companion, monthly billing. Repository catalog lists TradeFlowKit $149/month, TechDeck $99/month and PulseDesk $149/month; extra companions $29/month and seats $15/month. These are catalog references, not a verified live quote. On 2026-10-05 the public pricing page and rendered browser exposed Price unavailable for the dynamic totals. Do not quote an unverified total or promise a free pilot. Confirm commercial terms and final Stripe Checkout before purchase. Source: https://operatoros.net/pricing',source=['packages/sdk/src/products.ts','https://operatoros.net/pricing'])
    labels=[('1',d['key'],'Your flagship application'),('5','Team seats','Included with the stack'),('1','Eligible companion','Chosen for your organization')]
    for i,(num,a,c) in enumerate(labels):
        x=48+i*296;b.text(num,x,174,274,100,p['accent'],'Light');b.text(a,x,297,274,25,font='Bold');b.text(c,x,342,274,16,p['muted'])
    b.band('Free OperatorOS home base. Monthly billing. Add seats and eligible companions as needed.',404,47)
    b.text('Confirm the current quote and any usage charges before activation.',48,465,626,13,p['muted'])
    b.text('Review plans →',737,463,175,15,p['accent'],'Bold',align='right',url='https://operatoros.net/pricing?product='+d['key'].lower()+'#build-stack')
    # 16: Concrete next action.
    b.new('Prove the fit\nwith a focused 30-day pilot.','PROPOSED EVALUATION / SCOPE AND TERMS BY AGREEMENT',notes='This is a proposed evaluation structure, not an existing free-trial or onboarding-service promise. Agree pricing, support, scope, data handling, responsible owners and success criteria before starting. Use comparable baseline and pilot samples. '+d['pilot_exit'])
    for i,((tag,c),period) in enumerate(zip(d['pilot'],['WEEK 1','WEEKS 2–3','WEEK 4'])):
        x=48+i*296;b.label(period,x,196,272);b.line(x,225,x+272,225,p['accent'],2)
        b.text(tag,x,249,272,20,font='Bold',maxh=52);b.text(c,x,311,265,19,p['muted'],maxh=118)
    b.text(d['pilot_exit'],48,439,864,17.5,p['accent'],'Bold',maxh=55)
    # 17: Powerful buyer-centered close. End main presentation here.
    b.new(d['close_title'],'LET’S MAKE THE NEXT HANDOFF BETTER',notes='CLOSE: Ask for one representative example, a workflow owner and a success measure. Agree a scoped working session and the commercial next step. The CTA opens the existing OperatorOS contact page; no appointment or service commitment is automatically created.',cover=True)
    b.picture(ROOT/'assets'/d['image'],650,0,310,H);b.rect(630,0,20,H,p['accent']);b.brand()
    b.label('THE NEXT CONVERSATION',48,116)
    b.text(d['close_title'],48,158,570,44,font='Bold',maxh=165)
    b.text(d['close_copy'],48,339,535,20,p['muted'],maxh=96)
    b.rect(48,445,321,43,p['accent'],r=4);b.text('Map your workflow with us →',64,456,293,16,p['bg'],'Bold',url='https://operatoros.net/john')
    b.text('operatoros.net/for/'+d['lane'],48,509,562,12,p['muted'],url='https://operatoros.net/for/'+d['lane'])
    # 18: Objection backup.
    b.new('The questions that\nmake the fit clearer.','APPENDIX / BUYER QUESTIONS',notes='Use only the relevant objection. Invite a specific example from the buyer instead of treating this as a scripted rebuttal. Unsupported integrations or clinical capabilities are material fit considerations.')
    for i,(a,c) in enumerate(d['objections']):
        y=181+i*100;b.line(48,y,912,y);b.text(a,48,y+14,302,19,font='Bold',maxh=74);b.text(c,389,y+14,523,17,p['muted'],maxh=82)
    # 19: Detailed appendix capability list.
    b.new('Capabilities at a glance.','APPENDIX / FEATURE CHECKLIST',notes='This list is a map of source-supported workflow families, not proof that every legacy feature or external provider is production accepted. Verify the buyer’s precise required workflow and role on the target tenant. Details: '+', '.join(d['extra_sources']))
    for i,(a,c) in enumerate(d['checklist']):
        x=48+(i%2)*444;y=148+(i//2)*56;b.line(x,y+52,x+420,y+52);b.text(a,x,y+4,420,17,p['accent'],'Bold');b.text(c,x,y+28,415,14.2,p['muted'])
    # 20: Shareable scope and source references, not internal secrets.
    b.new('Agree the scope.\nVerify the workflow.','APPENDIX / EVALUATION NOTES',notes='Internal release context for presenter: current local source '+COMMIT+' includes the OpenAI enablement merge. The recorded published release gate remains v65, with semantic provider disabled; the later hosted OpenAI credential recovery/publication and authenticated tenant acceptance remain open in implementation status. Do not sell AI activation as universal. This deck is collateral, not release acceptance. '+d['scope'][2][1],source=COMMON_SOURCES+d['extra_sources'])
    for i,(a,c) in enumerate(d['scope']):
        y=178+i*73;b.text(a,48,y,226,18,p['accent'],'Bold');b.text(c,304,y,608,17,p['muted'],maxh=66)
    b.line(48,408,912,408)
    b.text('Scenario examples are illustrative. Workflow diagrams are explanatory, not product screenshots.',48,423,864,11.5,p['muted'])
    b.text('Product scope reviewed 05 Oct 2026. Validate the selected tenant and connected services before rollout.',48,445,864,11.5,p['muted'])
    b.text('Explore this lane →',48,472,290,12,p['accent'],'Bold',url='https://operatoros.net/for/'+d['lane'])
    b.text('Plans and package details →',377,472,319,12,p['accent'],'Bold',url='https://operatoros.net/pricing')
    b.text('Product help →',750,472,162,12,p['accent'],'Bold',align='right',url='https://operatoros.net/help')
    return b.save()

if __name__=='__main__':
    manifests=[make(d) for d in DECKS]
    (ROOT/'slide_index.json').write_text(json.dumps(manifests,indent=2),encoding='utf-8')
    print(json.dumps([{'product':m['product'],'slides':m['slides'],'pdf':m['pdf'],'pptx':m['pptx']} for m in manifests],indent=2))
