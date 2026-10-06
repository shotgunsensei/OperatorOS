from pathlib import Path
import json, re, zipfile, hashlib
from collections import defaultdict
import xml.etree.ElementTree as ET
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader
import pypdfium2 as pdfium
from pptx import Presentation

ROOT=Path(__file__).resolve().parent
PREVIEW=ROOT/'previews';PREVIEW.mkdir(exist_ok=True)
results=[]
for manifest in json.loads((ROOT/'slide_index.json').read_text()):
    stem=Path(manifest['pdf']).stem
    pdfpath=ROOT/manifest['pdf'];pptxpath=ROOT/manifest['pptx']
    pdf=PdfReader(pdfpath);prs=Presentation(pptxpath)
    problems=[];overlaps=[];slide_elements=defaultdict(list)
    boxes=json.loads((ROOT/(stem+'_Layout.json')).read_text())
    for box in boxes:slide_elements[box['slide']].append(box)
    if len(pdf.pages)!=20 or len(prs.slides)!=20:problems.append('Wrong page/slide count')
    for n,slide in enumerate(prs.slides,1):
        if len(slide.notes_slide.notes_text_frame.text)<100:problems.append(f'Missing notes on {n}')
        for sh in slide.shapes:
            if sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width+12700 or sh.top+sh.height>prs.slide_height+12700:
                problems.append(f'PowerPoint shape beyond page on {n}')
        page=pdf.pages[n-1]
        if tuple(float(v) for v in page.mediabox[2:])!=(960,540):problems.append(f'Wrong PDF size on {n}')
        actual=re.sub(r'\s+',' ',page.extract_text()).strip()
        for box in slide_elements[n]:
            expected=re.sub(r'\s+',' ',box['text']).strip()
            if expected not in actual:problems.append(f'PDF text mismatch on {n}: {expected[:50]}')
        for i,a in enumerate(slide_elements[n]):
            ax,ay,aw,ah=a['box']
            for b in slide_elements[n][i+1:]:
                bx,by,bw,bh=b['box']
                ix=min(ax+aw,bx+bw)-max(ax,bx);iy=min(ay+ah,by+bh)-max(ay,by)
                if ix>3 and iy>3:overlaps.append(dict(slide=n,a=a['text'],b=b['text'],intersection=[ix,iy]))
    font_status={}
    for page in pdf.pages:
        for k,v in page['/Resources']['/Font'].items():
            f=v.get_object();base=str(f.get('/BaseFont'));fd=f.get('/FontDescriptor')
            if fd:font_status[base]=any(t in fd.get_object() for t in ['/FontFile','/FontFile2','/FontFile3'])
            elif base!='/Helvetica':font_status[base]=False
    with zipfile.ZipFile(pptxpath) as z:
        xmlfiles=[s for s in z.namelist() if s.endswith('.xml') or s.endswith('.rels')]
        for name in xmlfiles: ET.fromstring(z.read(name))
        slide_xmls=[n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+.xml',n)]
        notes_xmls=[n for n in z.namelist() if re.fullmatch(r'ppt/notesSlides/notesSlide\d+.xml',n)]
        if len(notes_xmls)!=20:problems.append('Incorrect notes count')
    doc=pdfium.PdfDocument(str(pdfpath));thumbs=[]
    for n,page in enumerate(doc):
        im=page.render(scale=1.4).to_pil().convert('RGB')
        im.save(PREVIEW/f'{manifest["product"]}_{n+1:02d}.png')
        im.thumbnail((480,270));thumbs.append(im.copy())
    sheet=Image.new('RGB',(4*500,5*302),'#D7DFE5');draw=ImageDraw.Draw(sheet)
    font=ImageFont.truetype('C:/Windows/Fonts/segoeui.ttf',14)
    for i,im in enumerate(thumbs):
        x=(i%4)*500+10;y=(i//4)*302+8;sheet.paste(im,(x,y));draw.text((x,y+274),f'{i+1:02d} / {manifest["index"][i]["title"][:55]}',font=font,fill='#19364B')
    sheet.save(PREVIEW/f'{manifest["product"]}_Contact_Sheet.jpg',quality=92)
    results.append(dict(product=manifest['product'],pages=len(pdf.pages),slides=len(prs.slides),notes=len(notes_xmls),
        editable_text_blocks=len(boxes),rendered_pages=len(thumbs),problems=problems,text_overlap_candidates=overlaps,
        pdf_embedded_fonts=font_status,pdf_links=sum(len(p.get('/Annots',[])) for p in pdf.pages),
        pptx_bytes=pptxpath.stat().st_size,pdf_bytes=pdfpath.stat().st_size,
        pdf_sha256=hashlib.sha256(pdfpath.read_bytes()).hexdigest(),pptx_sha256=hashlib.sha256(pptxpath.read_bytes()).hexdigest()))
(ROOT/'quality_check.json').write_text(json.dumps(results,indent=2),encoding='utf-8')
print(json.dumps(results,indent=2))
